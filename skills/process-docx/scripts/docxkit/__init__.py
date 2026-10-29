"""Public API for the process-docx skill.

Importing this package bootstraps python-docx into the local ``scripts/.deps``
directory when it is not already available in the current environment.
"""

from __future__ import annotations

from ._bootstrap import DEPS_DIR, doctor, ensure_deps

ensure_deps()

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .core import (
    add_heading,
    add_page_break,
    add_paragraph,
    add_table,
    delete_paragraph,
    duplicate_paragraph,
    extract_images,
    extract_text,
    insert_paragraph_after,
    iter_blocks,
    list_styles,
    new_docx,
    open_docx,
    read_table,
    remove_comments,
    replace_text,
    save_docx,
    set_footer_text,
    set_header_text,
    set_page_margins,
    set_page_size,
    set_paragraph_spacing,
    set_run_font,
    write_table,
)
from .inspection import inspect_docx
from .ooxml import (
    dump_xml,
    paragraph_list_level,
    set_cell_shading,
    set_paragraph_border,
    set_paragraph_shading,
)

__version__ = "0.1.0"

__all__ = [
    "DEPS_DIR",
    "Document",
    "Paragraph",
    "Run",
    "Table",
    "add_heading",
    "add_page_break",
    "add_paragraph",
    "add_table",
    "delete_paragraph",
    "doctor",
    "dump_xml",
    "duplicate_paragraph",
    "ensure_deps",
    "extract_images",
    "extract_text",
    "insert_paragraph_after",
    "inspect_docx",
    "iter_blocks",
    "list_styles",
    "new_docx",
    "open_docx",
    "paragraph_list_level",
    "read_table",
    "remove_comments",
    "replace_text",
    "save_docx",
    "set_cell_shading",
    "set_footer_text",
    "set_header_text",
    "set_page_margins",
    "set_page_size",
    "set_paragraph_border",
    "set_paragraph_shading",
    "set_paragraph_spacing",
    "set_run_font",
    "write_table",
]
